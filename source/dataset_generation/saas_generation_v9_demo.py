"""
saas_generation_v9_demo.py

DEMO ONLY
---------
This demonstrates how the code runs. The demo tables it produces are not the
final ones used in the repo. The final tables are too large to host here.
Contact the author if you need a demo of the actual code.

PURPOSE
-------
Generates a small, synthetic, long-format "customer-month panel" that mimics
a SaaS subscription history. It also writes a cost-per-customer table and an
Excel workbook of helper tables.

HOW IT WORKS
------------
All constants and growth assumptions are defined in calibration_config_demo.py. 
Since it is just a demo, curves are flat and every channel and plan behaves
about the same. Results from this demo config do not resemble the final
curves, ratios or prices.  Do not expect the data generated here to produce
any meaningful insights. 

The engine code is the same as the full version. Only the config differs.

OUTPUT
------
- saas_dataset_v9_demo_<span>_<timestamp>.csv
- cost_per_customer_v9_demo_<span>_<timestamp>.csv
- helper_dfs/helper_dfs_v9_demo_<span>_<timestamp>.xlsx

Panel columns include customer_clv_24m, customer_cac,
unit_clv_cac_ratio_24m and is_mature_24m. To aggregate CLV, filter
is_mature_24m = TRUE and lifetime_month = 0 (one row per customer).
"""

import os
import datetime

import numpy as np
import pandas as pd


# ======================================================================
# SECTION 1: CONFIGURATION (DEMO)
# ======================================================================
# Every constant comes from calibration_config_demo.py. This file has no
# defaults of its own. If the import fails, the script stops here.

try:
    from calibration_config_demo import (
        SEED,
        DEFAULT_START_MONTH,
        DEFAULT_NUM_MONTHS,
        CLV_HORIZON_MONTHS,
        GROWTH_CAP_MONTHS,
        BASE_SIGNUPS,
        SIGNUP_GROWTH,
        SIGNUP_NOISE,
        COHORT_RET_NOISE,
        MRR_NOISE,
        CONTRACTION_PROB,
        CONTRACTION_RANGE,
        CAC_NOISE,
        CHANNELS,
        PLANS,
    )
except ImportError as exc:
    raise RuntimeError(
        "\n\ncalibration_config_demo.py not found.\n\n"
        "Place it in the same folder as this script.\n"
    ) from exc

DEMO_NOTICE = (
    "DEMO only. This demonstrates how the code runs. "
    "The produced demo tables are not the final ones used in the repo. "
    "The final tables are too large. "
    "Contact the author if you need a demo of the actual code."
)


def print_demo_notice():
    border = "=" * 78
    print(border)
    print(DEMO_NOTICE)
    print(border)


OUT_DIR = os.path.dirname(os.path.abspath(__file__))


# ======================================================================
# SECTION 2: COHORT CALENDAR HELPERS
# ======================================================================

def build_calendar(start_month, num_months):
    cohort_list = []
    first_period = pd.Period(start_month, freq="M")
    for offset in range(num_months):
        cohort_list.append(first_period + offset)
    return cohort_list


def month_diff(start_period, end_period):
    return (end_period.year - start_period.year) * 12 + (end_period.month - start_period.month)


def resolve_horizon(start_month=None, num_months=None, start_year=None, end_year=None):
    """
    Resolves the (start_month, num_months) pair actually used to build the
    calendar, from whichever inputs the caller provided:

    - num_months given directly -> used as-is (with start_month, or the
      configured default start month if none given).
    - start_year + end_year given (num_months not given) -> num_months is
      computed as a whole-year span, e.g. 2022-2028 inclusive = 84 months.
    - Nothing given -> falls back to DEFAULT_START_MONTH / DEFAULT_NUM_MONTHS
      from calibration_config_demo.py.
    """
    if num_months is not None:
        resolved_start = start_month or DEFAULT_START_MONTH
        return resolved_start, num_months

    if start_year is not None and end_year is not None:
        if end_year < start_year:
            raise ValueError(f"end_year ({end_year}) must be >= start_year ({start_year})")
        resolved_start = f"{start_year}-01"
        resolved_num_months = (end_year - start_year + 1) * 12
        return resolved_start, resolved_num_months

    return (start_month or DEFAULT_START_MONTH), DEFAULT_NUM_MONTHS


# ======================================================================
# SECTION 3: SIGNUP VOLUME HELPERS
# ======================================================================

def get_signup_volume(cohort_idx, base_signups, growth_rate, noise_pct, rng,
                       growth_cap_months=GROWTH_CAP_MONTHS):
    capped_idx = min(cohort_idx, growth_cap_months - 1)
    trended_count = base_signups * ((1 + growth_rate) ** capped_idx)
    noise_mult = 1 + rng.uniform(-noise_pct, noise_pct)
    return max(1, round(trended_count * noise_mult))


def build_signup_table(cohorts, base_signups, growth_rate, noise_pct, rng,
                        growth_cap_months=GROWTH_CAP_MONTHS):
    signup_table = {}
    for idx, cohort in enumerate(cohorts):
        signup_table[cohort] = get_signup_volume(
            idx, base_signups, growth_rate, noise_pct, rng, growth_cap_months
        )
    return signup_table


# ======================================================================
# SECTION 4: CHANNEL / PLAN-TIER MIX AND SEGMENT SPLITTING
# ======================================================================

def get_mix_share(config, cohort_idx, total_cohorts):
    start = config["mix_start"]
    end = config["mix_end"]
    if total_cohorts <= 1:
        return start
    progress = cohort_idx / (total_cohorts - 1)
    return start + (end - start) * progress


def split_signups(total_signups, cohort_idx, total_cohorts, channels, plans):
    splits = {}
    for channel, ch_cfg in channels.items():
        ch_share = get_mix_share(ch_cfg, cohort_idx, total_cohorts)
        for plan, pl_cfg in plans.items():
            pl_share = get_mix_share(pl_cfg, cohort_idx, total_cohorts)
            splits[(channel, plan)] = round(total_signups * (ch_share * pl_share))
    return splits


# ======================================================================
# SECTION 5: RETENTION & EXPANSION CURVE HELPERS
# ======================================================================

def interpolate_curve(anchors, max_months):
    """
    Interpolates between anchor points across the requested horizon.
    np.interp holds the curve flat at the value of the last anchor key for
    any timeline point beyond it. This is what keeps long-tracked cohorts
    stabilized at a plateau (e.g. 45% retention) instead of decaying toward
    zero, for ANY max_months, not just the anchor points defined above.
    """
    months = sorted(anchors.keys())
    values = [anchors[m] for m in months]
    timeline = np.arange(max_months)
    return np.interp(timeline, months, values)


def apply_plan_modifier(base_curve, modifier_pts):
    """
    Shift the retention curve up or down by the plan's modifier (in points).
    Month 0 is always pinned to 100%. Everyone who signs up pays in their
    first month, so a plan shift only applies from month 1 onward.
    """
    modified = base_curve + (modifier_pts / 100.0)
    modified = np.clip(modified, 0.0, 1.0)
    modified[0] = 1.0
    return modified


def apply_retention_noise(base_curve, max_noise_pts, rng):
    """
    Add one random shift (in points) to the whole retention curve.
    Month 0 stays pinned at 100%. The random draw is unchanged.
    """
    noise = rng.uniform(-max_noise_pts / 100.0, max_noise_pts / 100.0)
    noisy = np.clip(base_curve + noise, 0.0, 1.0)
    noisy[0] = 1.0
    return noisy


def get_customer_counts(initial_count, retention_curve):
    counts = np.round(retention_curve * initial_count)
    non_increasing = np.minimum.accumulate(counts)
    non_increasing[0] = initial_count
    return [int(val) for val in non_increasing]


def get_departures(initial_count, active_counts):
    departures = [None] * initial_count
    idx = 0
    for month in range(1, len(active_counts)):
        dropped = active_counts[month - 1] - active_counts[month]
        for _ in range(dropped):
            departures[idx] = month - 1
            idx += 1
    return departures


def apply_expansion_scaler(expansion_curve, scaler):
    """
    Scale how much a channel's customers grow their revenue.

    The plan's expansion curve starts at 1.00 (what a customer paid in
    month 0). A value of 1.35 means they now pay 35% more, from extra seats,
    add-ons, or usage. This function scales only the growth above 1.00.
    A scaler of 1.20 makes growth 20% bigger. A scaler of 0.92 makes it 8%
    smaller. The result never drops below 1.00.
    """
    growth_portion = (expansion_curve - 1.0) * scaler
    return np.maximum(1.0 + growth_portion, 1.0)


# ======================================================================
# SECTION 5B: MATHEMATICAL SOLVER (NO CALIBRATION BUFFER)
# ======================================================================

def get_expected_billing_units(channel, plan, num_months):
    """
    How many full-price months one new customer is expected to pay over
    num_months. One unit = one month billed at the plan's base price.

    Each month is: share still paying x revenue growth x downgrade factor.
    The result is added up over the horizon. It uses the smooth curves, with
    no random noise. It feeds only the expected-ratio column in the helper
    workbook. It does not set any price.
    """
    # Settings for this channel (Ads, Organic or Referral).
    ch_cfg = CHANNELS[channel]
    # Settings for this plan (Basic, Pro or Enterprise).
    pl_cfg = PLANS[plan]

    # Share of customers still paying in each month, from the channel's anchor points.
    ret_curve = interpolate_curve(ch_cfg["retention_anchors"], num_months)
    # Shift that share up or down by the plan's retention modifier (month 0 stays 100%).
    ret_curve = apply_plan_modifier(ret_curve, pl_cfg["retention_modifier"])

    # Revenue growth of a staying customer, from the plan's anchor points (1.00 = month-0 revenue).
    exp_curve = interpolate_curve(pl_cfg["expansion_anchors"], num_months)
    # Scale the growth part by the channel's expansion scaler.
    exp_curve = apply_expansion_scaler(exp_curve, ch_cfg["expansion_scaler"])

    # Average size of a random downgrade. The range is 10% to 30%, so the average is 20%.
    avg_contraction_size = sum(CONTRACTION_RANGE) / 2
    # Revenue kept per month after downgrades: 1 - (5% chance x 20% size) = 0.99.
    # This is the same every month because each month's downgrade is drawn fresh.
    contraction_factor = 1.0 - (CONTRACTION_PROB * avg_contraction_size)

    # Running total of billing units.
    expected_billing_units = 0.0
    # Go through each month of the horizon (0 to 23 for a 24-month horizon).
    for t in range(num_months):
        # Share still paying x revenue growth x downgrade factor, added to the total.
        expected_billing_units += ret_curve[t] * exp_curve[t] * contraction_factor

    return expected_billing_units


def get_expected_clv_cac_ratio(channel, plan, num_months):
    """
    Expected CLV:CAC over num_months for one channel and plan.
        (plan base price x expected billing units) / starting CAC
    Used to cross-check the realized ratios in the panel.
    """
    units = get_expected_billing_units(channel, plan, num_months)
    cac = CHANNELS[channel]["base_cac"] * PLANS[plan]["cac_mult"]
    return round(PLANS[plan]["base_arpu"] * units / cac, 2)


# ======================================================================
# SECTION 6: REVENUE HELPERS
# ======================================================================

def get_mrr(base_arpu, expansion_factor, noise_pct, contraction_prob, contraction_range, rng):
    """
    One customer's revenue for one month.
    base ARPU x revenue growth factor for this lifetime month x a random
    +/- noise. Then a small chance of a one-month downgrade. The downgrade
    is not remembered. Next month starts fresh from the base.
    """
    noise = 1 + rng.uniform(-noise_pct, noise_pct)
    revenue = base_arpu * expansion_factor * noise

    if rng.random() < contraction_prob:
        contraction = 1 - rng.uniform(*contraction_range)
        revenue *= contraction

    return round(revenue, 2)


def get_segment_rows(cohort, channel, plan, initial_count, departures, arpu, expansion_curve,
                     noise_pct, contraction_prob, contraction_range, max_observable, prefix, rng):
    rows = []
    for customer_idx in range(initial_count):
        customer_id = f"{prefix}_{customer_idx:05d}"
        last_active = departures[customer_idx]
        if last_active is None:
            last_active = max_observable

        for t in range(last_active + 1):
            rows.append({
                "customer_id": customer_id,
                "channel": channel,
                "plan_tier": plan,
                "cohort_month": str(cohort),
                "calendar_month": str(cohort + t),
                "lifetime_month": t,
                "is_active": True,
                "mrr": get_mrr(arpu, expansion_curve[t], noise_pct, contraction_prob, contraction_range, rng),
            })
    return rows


# ======================================================================
# SECTION 7: CAC / COST-PER-CUSTOMER HELPERS
# ======================================================================

def get_cac(ch_cfg, plan_mult, cohort_idx, rng, growth_cap_months=GROWTH_CAP_MONTHS):
    capped_idx = min(cohort_idx, growth_cap_months - 1)
    base = ch_cfg["base_cac"] * plan_mult
    trended = base * ((1 + ch_cfg["cac_trend"]) ** capped_idx)
    noise = 1 + rng.uniform(-CAC_NOISE, CAC_NOISE)
    return round(trended * noise, 2)


def build_cac_table(cohorts, channels, plans, new_customers_lookup, rng,
                     growth_cap_months=GROWTH_CAP_MONTHS):
    records = []
    total_cohorts = len(cohorts)
    for idx, cohort in enumerate(cohorts):
        for channel, ch_cfg in channels.items():
            ch_share = get_mix_share(ch_cfg, idx, total_cohorts)
            for plan, pl_cfg in plans.items():
                pl_share = get_mix_share(pl_cfg, idx, total_cohorts)
                mix_share = ch_share * pl_share

                new_custs = new_customers_lookup.get((cohort, channel, plan), 0)
                cac = get_cac(ch_cfg, pl_cfg["cac_mult"], idx, rng, growth_cap_months)
                total_cost = round(new_custs * cac, 2)

                records.append({
                    "cohort_month": str(cohort),
                    "channel": channel,
                    "plan_tier": plan,
                    "new_customers": new_custs,
                    "segment_mix_share_percentage": round(mix_share, 4),
                    "cac_per_customer_dollars": cac,
                    "total_acquisition_cost_dollars": total_cost,
                })
    return records


# ======================================================================
# SECTION 8: OUTPUT / FILE-SAVING HELPERS
# ======================================================================

def get_timestamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


def save_panel(df, out_dir, timestamp, span_tag):
    path = os.path.join(out_dir, f"saas_dataset_v9_demo_{span_tag}_{timestamp}.csv")
    df.to_csv(path, index=False)
    return path


def save_cac_table(df, out_dir, timestamp, span_tag):
    path = os.path.join(out_dir, f"cost_per_customer_v9_demo_{span_tag}_{timestamp}.csv")
    df.to_csv(path, index=False)
    return path


def build_helper_dataframes(cohorts, num_months, max_lifetime, signup_table, new_customers_lookup):
    """
    Builds the intermediate structures used during generation (signup
    volumes, retention/expansion curves, plan pricing and expected ratios, segment mix shares,
    new-customer counts) into standalone DataFrames, purely for my inspection
    and debugging. These are not needed downstream by the panel or CAC
    table, which already have everything they need baked in.
    """
    helper_dfs = {}

    helper_dfs["signup_table"] = pd.DataFrame([
        {"cohort_month": str(cohort), "total_signups": signup_table[cohort]}
        for cohort in cohorts
    ])

    retention_rows = []
    expansion_rows = []
    price_rows = []
    for channel, ch_cfg in CHANNELS.items():
        base_ret = interpolate_curve(ch_cfg["retention_anchors"], max_lifetime)
        for plan, pl_cfg in PLANS.items():
            plan_ret = apply_plan_modifier(base_ret, pl_cfg["retention_modifier"])
            for lifetime_month, retention_rate in enumerate(plan_ret):
                retention_rows.append({
                    "channel": channel,
                    "plan_tier": plan,
                    "lifetime_month": lifetime_month,
                    "retention_rate": round(float(retention_rate), 4),
                })

            plan_exp = interpolate_curve(pl_cfg["expansion_anchors"], max_lifetime)
            segment_exp = apply_expansion_scaler(plan_exp, ch_cfg["expansion_scaler"])
            for lifetime_month, expansion_multiplier in enumerate(segment_exp):
                expansion_rows.append({
                    "channel": channel,
                    "plan_tier": plan,
                    "lifetime_month": lifetime_month,
                    "expansion_multiplier": round(float(expansion_multiplier), 4),
                })

            units = get_expected_billing_units(channel, plan, CLV_HORIZON_MONTHS)
            price_rows.append({
                "channel": channel,
                "plan_tier": plan,
                "base_arpu": pl_cfg["base_arpu"],
                "expected_billing_units": round(units, 2),
                "base_cac_dollars": ch_cfg["base_cac"] * pl_cfg["cac_mult"],
                "clv_horizon_months": CLV_HORIZON_MONTHS,
                "expected_clv_cac_24m": get_expected_clv_cac_ratio(channel, plan, CLV_HORIZON_MONTHS),
            })

    helper_dfs["retention_curves"] = pd.DataFrame(retention_rows)
    helper_dfs["expansion_curves"] = pd.DataFrame(expansion_rows)
    helper_dfs["pricing_and_expected_ratio"] = pd.DataFrame(price_rows)

    channel_mix_rows = []
    plan_mix_rows = []
    for cohort_idx, cohort in enumerate(cohorts):
        for channel, ch_cfg in CHANNELS.items():
            channel_mix_rows.append({
                "cohort_month": str(cohort),
                "cohort_idx": cohort_idx,
                "channel": channel,
                "mix_share": round(get_mix_share(ch_cfg, cohort_idx, num_months), 4),
            })
        for plan, pl_cfg in PLANS.items():
            plan_mix_rows.append({
                "cohort_month": str(cohort),
                "cohort_idx": cohort_idx,
                "plan_tier": plan,
                "mix_share": round(get_mix_share(pl_cfg, cohort_idx, num_months), 4),
            })

    helper_dfs["channel_mix_shares"] = pd.DataFrame(channel_mix_rows)
    helper_dfs["plan_mix_shares"] = pd.DataFrame(plan_mix_rows)

    helper_dfs["new_customers_by_segment"] = pd.DataFrame([
        {
            "cohort_month": str(cohort),
            "channel": channel,
            "plan_tier": plan,
            "new_customers": new_customer_count,
        }
        for (cohort, channel, plan), new_customer_count in new_customers_lookup.items()
    ])

    return helper_dfs


HELPER_DESCRIPTIONS = {
    "signup_table": "Total signups per cohort month.",
    "channel_mix_shares": "Each channel's share of signups, by cohort month.",
    "plan_mix_shares": "Each plan's share of signups, by cohort month.",
    "new_customers_by_segment": "Customers acquired per cohort month, channel and plan.",
    "retention_curves": "Expected share of customers still paying, by lifetime month. Before cohort noise.",
    "expansion_curves": "Revenue growth multiplier by lifetime month. 1.00 = month-0 revenue.",
    "pricing_and_expected_ratio": "Plan list price, expected billing months, starting CAC, and the expected 24-month CLV:CAC for each channel and plan.",
}


def save_helper_dfs(helper_dfs, out_dir, timestamp, span_tag):
    """
    Writes every helper table into ONE Excel workbook, one sheet per table.
    A first sheet called "contents" lists the sheets, their row counts, and
    what each one holds. Values only, no formulas.
    """
    from openpyxl.styles import Font, PatternFill

    helper_dir = os.path.join(out_dir, "helper_dfs")
    os.makedirs(helper_dir, exist_ok=True)
    path = os.path.join(helper_dir, f"helper_dfs_v9_demo_{span_tag}_{timestamp}.xlsx")

    contents = pd.DataFrame([
        {"sheet": name, "rows": len(df), "description": HELPER_DESCRIPTIONS.get(name, "")}
        for name, df in helper_dfs.items()
    ])

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        sheets = {"contents": contents, **helper_dfs}
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            for row in ws.iter_rows():
                for cell in row:
                    cell.font = Font(name="Arial", size=10, bold=(cell.row == 1))
                    if cell.row == 1:
                        cell.fill = PatternFill("solid", start_color="D9E1F2")
            for col_cells in ws.columns:
                longest = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
                ws.column_dimensions[col_cells[0].column_letter].width = min(max(longest + 2, 10), 80)
    return path


# ======================================================================
# SECTION 9: MAIN ORCHESTRATION
# ======================================================================

def main(start_month=None, num_months=None, start_year=None, end_year=None):
    """
    Generates the panel. All time horizon inputs are optional:

    - No arguments -> falls back to calibration_config_demo.py's
      DEFAULT_START_MONTH / DEFAULT_NUM_MONTHS.
    - num_months=N -> generates N cohorts starting at start_month (or the
      configured default start month if start_month is also omitted).
    - start_year=Y1, end_year=Y2 -> generates whole years Y1 through Y2
      inclusive, e.g. start_year=2022, end_year=2028 -> 84 months.

    Customer lifetime tracking (MAX_LIFETIME) always extends to match the
    full resolved time horizon.
    """
    print_demo_notice()
    resolved_start_month, resolved_num_months = resolve_horizon(
        start_month, num_months, start_year, end_year
    )
    num_months = resolved_num_months
    max_lifetime = resolved_num_months  # lifetime tracking = full span

    rng = np.random.default_rng(SEED)

    cohorts = build_calendar(resolved_start_month, num_months)
    signup_table = build_signup_table(
        cohorts, BASE_SIGNUPS, SIGNUP_GROWTH, SIGNUP_NOISE, rng, GROWTH_CAP_MONTHS
    )

    all_rows = []
    new_customers_lookup = {}
    final_cohort = cohorts[-1]
    span_tag = f"{cohorts[0]}_{final_cohort}"

    for idx, cohort in enumerate(cohorts):
        total_signups = signup_table[cohort]
        splits = split_signups(total_signups, idx, num_months, CHANNELS, PLANS)
        months_remaining = month_diff(cohort, final_cohort)
        max_observable = min(max_lifetime - 1, months_remaining)

        for (channel, plan), initial_count in splits.items():
            new_customers_lookup[(cohort, channel, plan)] = initial_count

            if initial_count <= 0:
                continue

            print(f"Generating segment: {cohort} | {channel} / {plan} ({initial_count} customers)")

            ch_cfg = CHANNELS[channel]
            pl_cfg = PLANS[plan]

            base_ret = interpolate_curve(ch_cfg["retention_anchors"], max_lifetime)
            plan_ret = apply_plan_modifier(base_ret, pl_cfg["retention_modifier"])
            noisy_ret = apply_retention_noise(plan_ret, COHORT_RET_NOISE, rng)

            obs_ret_slice = noisy_ret[: max_observable + 1]
            active_counts = get_customer_counts(initial_count, obs_ret_slice)
            departures = get_departures(initial_count, active_counts)

            # One list price per plan. Every channel pays the same base price.
            arpu = pl_cfg["base_arpu"]

            plan_exp = interpolate_curve(pl_cfg["expansion_anchors"], max_lifetime)
            segment_exp = apply_expansion_scaler(plan_exp, ch_cfg["expansion_scaler"])
            obs_exp_slice = segment_exp[: max_observable + 1]

            prefix = f"{cohort}_{channel}_{plan}"
            segment_rows = get_segment_rows(
                cohort, channel, plan, initial_count, departures, arpu, obs_exp_slice,
                MRR_NOISE, CONTRACTION_PROB, CONTRACTION_RANGE, max_observable, prefix, rng
            )
            all_rows.extend(segment_rows)

    panel_df = pd.DataFrame(all_rows)
    cac_records = build_cac_table(
        cohorts, CHANNELS, PLANS, new_customers_lookup, rng, GROWTH_CAP_MONTHS
    )
    cac_df = pd.DataFrame(cac_records)

    print("Pre-calculating unit economic fields to prevent Excel overhead...")

    # Step 1: Cap CLV at the first CLV_HORIZON_MONTHS lifetime months
    capped = panel_df[panel_df["lifetime_month"] < CLV_HORIZON_MONTHS]
    clv_map = capped.groupby("customer_id")["mrr"].sum().round(2).to_dict()
    panel_df["customer_clv_24m"] = panel_df["customer_id"].map(clv_map)

    # Step 2: Extract individual customer acquisition costs from the segment records
    cac_map = cac_df.set_index(["cohort_month", "channel", "plan_tier"])["cac_per_customer_dollars"].to_dict()

    # Step 3: Match segment indicators back to panel rows to compute ratios
    panel_df["customer_cac"] = panel_df.set_index(["cohort_month", "channel", "plan_tier"]).index.map(cac_map)
    panel_df["unit_clv_cac_ratio_24m"] = (panel_df["customer_clv_24m"] / panel_df["customer_cac"]).round(4)

    # Step 4: Flag cohorts observed for the full CLV horizon
    mature_map = {
        str(c): month_diff(c, final_cohort) >= CLV_HORIZON_MONTHS - 1 for c in cohorts
    }
    panel_df["is_mature_24m"] = panel_df["cohort_month"].map(mature_map)

    timestamp = get_timestamp()
    panel_path = save_panel(panel_df, OUT_DIR, timestamp, span_tag)
    cac_path = save_cac_table(cac_df, OUT_DIR, timestamp, span_tag)

    helper_dfs = build_helper_dataframes(
        cohorts, num_months, max_lifetime, signup_table, new_customers_lookup
    )
    helper_path = save_helper_dfs(helper_dfs, OUT_DIR, timestamp, span_tag)

    print(f"Horizon timeframe             : {resolved_start_month} + {num_months} months "
          f"({cohorts[0]} - {final_cohort})")
    print(f"Customer-month panel saved to : {panel_path}")
    print(f"  Rows generated              : {len(panel_df):,}")
    print(f"  Unique customers            : {panel_df['customer_id'].nunique():,}")
    print(f"  Columns added               : customer_clv_24m, customer_cac, unit_clv_cac_ratio_24m, is_mature_24m")
    print(f"Cost-per-customer table saved to: {cac_path}")
    print(f"  Rows generated              : {len(cac_df):,}")
    print(f"Helper workbook saved to      : {helper_path}")
    for name, df in helper_dfs.items():
        print(f"  sheet {name:28s}: {len(df):,} rows")

    print()
    print_demo_notice()
    return panel_df, cac_df


if __name__ == "__main__":
    # Small demo run. 36 monthly cohorts starting 2022-01.
    main(num_months=36)
